"""DOCX 接口文档解析器（重构版）

解析 Word 格式的接口文档，支持：
- Heading 标题层次结构
- 多种表格格式（3列/5列/9列等）
- 通过配置驱动的表头映射
"""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .base_parser import BaseParser, ParsedField, ParsedView, ParseResult
from .header_mapper import HeaderMapper
from .table_classifier import TableClassifier, TableType


class DocxParser(BaseParser):
    """DOCX 解析器"""

    def __init__(self, config_path: Optional[str] = None):
        self._mapper = HeaderMapper(config_path)
        self._classifier = TableClassifier(self._mapper)

    def parse(self, file_path: str) -> ParseResult:
        """解析 DOCX 文件"""
        path = Path(file_path)
        if not path.exists():
            return ParseResult(file_path=str(file_path), errors=[f"文件不存在: {file_path}"])

        try:
            from docx import Document
        except ImportError:
            return ParseResult(file_path=str(file_path), errors=["缺少 python-docx 库，请执行: pip install python-docx"])

        try:
            doc = Document(str(path))
        except Exception as e:
            return ParseResult(file_path=str(file_path), errors=[f"打开文件失败: {e}"])

        result = ParseResult(file_path=str(path))

        # 解析文档结构：标题 + 表格的关联
        views = self._parse_document_structure(doc)
        result.views = views

        return result

    def _parse_document_structure(self, doc) -> List[ParsedView]:
        """解析文档结构，将标题和表格关联"""
        views = []
        # 跟踪各级标题
        heading_stack = {}  # {level: text}
        # 跟踪最近的表名段落（用于数据采集类文档）
        last_table_name_para = ""

        # 遍历 body 中的所有元素
        for element in doc.element.body:
            tag = element.tag.split('}')[-1] if '}' in element.tag else element.tag

            if tag == 'p':
                # 段落 - 检查是否是标题
                heading_info = self._extract_heading(element)
                if heading_info:
                    level = heading_info['level']
                    text = heading_info['text']

                    # 更新当前级别标题
                    heading_stack[level] = text

                    # 清除更低级别的标题
                    for lower_level in list(heading_stack.keys()):
                        if lower_level > level:
                            del heading_stack[lower_level]
                else:
                    # 检查是否是包含表名的段落
                    para_text = self._get_paragraph_text(element)
                    if self._is_table_name_paragraph(para_text):
                        last_table_name_para = para_text

            elif tag == 'tbl':
                # 表格 - 尝试解析为字段定义表
                table_obj = self._find_table_object(doc, element)
                if table_obj is None:
                    continue

                # 优先使用包含ext_表名的段落，否则使用标题栈
                view_heading = ""
                if last_table_name_para:
                    view_heading = last_table_name_para
                    last_table_name_para = ""  # 重置，避免重复使用
                else:
                    # 使用最高级别的标题作为视图名称
                    for level in sorted(heading_stack.keys()):
                        if level <= 2:
                            view_heading = heading_stack[level]
                            break
                    if not view_heading and heading_stack:
                        view_heading = heading_stack[min(heading_stack.keys())]

                view = self._parse_table(table_obj, view_heading)
                if view:
                    views.append(view)

        return views

    def _get_paragraph_text(self, element) -> str:
        """获取段落纯文本"""
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        text_parts = element.findall('.//w:t', nsmap)
        return "".join(t.text or "" for t in text_parts).strip()

    def _is_table_name_paragraph(self, text: str) -> bool:
        """判断段落是否是表名段落

        支持的模式：
        1. 包含 ext_ 开头的标识符（如"扩展表-xxx ext_emr_activity_info"）
        2. 包含 V_ 开头的标识符（如"患者信息 V_PAT_INFO"）
        3. 包含下划线的英文标识符（如"挂号明细表 OUT_REGISTER"）
        4. 中文+空格+英文标识符模式
        """
        if not text:
            return False

        # 模式1: 包含 ext_ 开头的标识符
        if 'ext_' in text:
            return True

        # 模式2: 包含 V_ 开头的标识符
        if re.search(r'\bV_[A-Za-z_]\w+', text):
            return True

        # 模式3: 包含下划线的英文标识符（至少3个字符）
        if re.search(r'\b[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9_]+\b', text):
            return True

        # 模式4: 中文+空格+英文标识符
        if re.search(r'[一-龥]+\s+[A-Za-z_]\w+', text):
            return True

        return False

    def _extract_heading(self, element) -> Optional[Dict[str, Any]]:
        """从段落元素提取标题信息"""
        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}

        # 获取样式
        pstyle = element.find('.//w:pStyle', nsmap)
        if pstyle is None:
            return None

        style_val = pstyle.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '')

        # 检查是否是标题样式
        if not style_val.startswith('Heading') and not style_val.isdigit():
            return None

        # 提取级别
        level = 0
        if style_val.startswith('Heading'):
            try:
                level = int(style_val.replace('Heading', ''))
            except ValueError:
                level = 1
        elif style_val.isdigit():
            level = int(style_val)

        # 提取文本
        text_parts = element.findall('.//w:t', nsmap)
        full_text = "".join(t.text or "" for t in text_parts).strip()

        if not full_text:
            return None

        return {'text': full_text, 'level': level}

    def _find_table_object(self, doc, element):
        """根据 element 找到对应的 doc.tables 对象"""
        for table in doc.tables:
            if table._element is element:
                return table
        return None

    def _parse_table(self, table, heading_text: str) -> Optional[ParsedView]:
        """解析单个表格"""
        if len(table.rows) < 2:
            return None

        # 提取表头
        header_row = [cell.text.strip() for cell in table.rows[0].cells]

        # 提取数据行
        data_rows = []
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            if any(c for c in cells):
                data_rows.append(cells)

        if not data_rows:
            return None

        # 映射表头
        mapping = self._mapper.map_header(header_row)
        mapping = self._mapper.resolve_ambiguity(mapping, data_rows)

        # 分类表格
        classification = self._classifier.classify(header_row, data_rows)
        classification.header_mapping = mapping

        if not classification.is_field_table:
            return None

        # 提取视图标识
        view_code, view_name = self._extract_view_id(heading_text)

        # 提取字段
        fields = self._extract_fields(data_rows, mapping)

        if not fields:
            return None

        return ParsedView(
            view_code=view_code,
            view_name=view_name,
            fields=fields,
        )

    def _extract_view_id(self, heading_text: str) -> Tuple[str, str]:
        """从标题文本提取视图标识

        Returns:
            (view_code, view_name)
        """
        # 模式1: 中文名（英文名）
        match = re.search(r'([一-龥]+)[（(]([A-Za-z_]\w+)[)）]', heading_text)
        if match:
            return match.group(2), match.group(1)

        # 模式2: 中文名 ext_xxx（数据采集表格式）
        match = re.search(r'([一-龥][一-龥\-\s（）()]+)\s+(ext_\w+)', heading_text)
        if match:
            return match.group(2), match.group(1).strip()

        # 模式3: 纯英文标识符
        match = re.search(r'([A-Z][A-Za-z_]\w+)', heading_text)
        if match:
            return match.group(1), heading_text

        # 模式4: ext_开头的表名
        match = re.search(r'(ext_\w+)', heading_text)
        if match:
            return match.group(1), heading_text

        # 模式5: 纯中文标题
        # 去掉数字前缀
        clean_name = re.sub(r'^\d+(\.\d+)*\s*', '', heading_text).strip()
        return clean_name, clean_name

    def _extract_fields(self, data_rows: List[List[str]], mapping: Dict[int, str]) -> List[ParsedField]:
        """从数据行提取字段定义"""
        fields = []

        # 找到各字段的列索引
        name_col = None
        code_col = None
        type_col = None
        length_col = None
        required_col = None
        desc_col = None

        for col_idx, field_type in mapping.items():
            if field_type == "field_name":
                code_col = col_idx
            elif field_type == "field_chinese_name":
                name_col = col_idx
            elif field_type == "data_type":
                type_col = col_idx
            elif field_type == "data_length":
                length_col = col_idx
            elif field_type == "required":
                required_col = col_idx
            elif field_type == "description":
                desc_col = col_idx

        # 至少需要 code 或 name
        if code_col is None and name_col is None:
            return []

        for row in data_rows:
            # 获取字段代码
            code = ""
            if code_col is not None and code_col < len(row):
                code = row[code_col].strip().replace(" ", "")  # 去掉所有空格

            # 获取字段中文名
            name = ""
            if name_col is not None and name_col < len(row):
                name = row[name_col].strip()

            # 跳过空行
            if not code and not name:
                continue

            # 跳过明显不是字段定义的行
            if code in ["合计", "总计", "小计"] or name in ["合计", "总计", "小计"]:
                continue

            # 获取数据类型
            field_type = ""
            if type_col is not None and type_col < len(row):
                field_type = row[type_col].strip()

            # 获取长度
            length = 0
            if length_col is not None and length_col < len(row):
                length_str = row[length_col].strip()
                match = re.search(r'\d+', length_str)
                if match:
                    length = int(match.group())

            # 获取是否必填
            required = False
            if required_col is not None and required_col < len(row):
                req_str = row[required_col].strip().upper()
                required = req_str in ["Y", "YES", "是", "必填", "非空", "TRUE", "1"]

            # 获取备注
            description = ""
            if desc_col is not None and desc_col < len(row):
                description = row[desc_col].strip()

            fields.append(ParsedField(
                name=name,
                code=code,
                field_type=field_type,
                length=length,
                required=required,
                description=description,
            ))

        return fields


def parse_docx(file_path: str, config_path: Optional[str] = None) -> ParseResult:
    """便捷函数：解析 DOCX 文件"""
    parser = DocxParser(config_path)
    return parser.parse(file_path)
