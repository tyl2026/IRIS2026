"""XLSX 接口文档解析器

解析 Excel 格式的接口文档，支持：
- 多 Sheet 结构（每个 Sheet 是一个数据表）
- 自动识别表头行和数据行
- 通过配置驱动的表头映射
"""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .base_parser import BaseParser, ParsedField, ParsedView, ParseResult
from .header_mapper import HeaderMapper
from .table_classifier import TableClassifier, TableType


class XlsxParser(BaseParser):
    """XLSX 解析器"""

    def __init__(self, config_path: Optional[str] = None):
        self._mapper = HeaderMapper(config_path)
        self._classifier = TableClassifier(self._mapper)

    def parse(self, file_path: str) -> ParseResult:
        """解析 XLSX 文件"""
        path = Path(file_path)
        if not path.exists():
            return ParseResult(file_path=str(file_path), errors=[f"文件不存在: {file_path}"])

        try:
            import openpyxl
        except ImportError:
            return ParseResult(file_path=str(file_path), errors=["缺少 openpyxl 库，请执行: pip install openpyxl"])

        try:
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        except Exception as e:
            return ParseResult(file_path=str(file_path), errors=[f"打开文件失败: {e}"])

        result = ParseResult(file_path=str(path))

        try:
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                view = self._parse_sheet(ws, sheet_name)
                if view:
                    result.views.append(view)
        finally:
            wb.close()

        return result

    def _parse_sheet(self, ws, sheet_name: str) -> Optional[ParsedView]:
        """解析单个 Sheet"""
        # 读取所有行
        all_rows = []
        for row in ws.iter_rows(values_only=True):
            all_rows.append(row)

        if len(all_rows) < 3:  # 至少需要标题行 + 表头行 + 1行数据
            return None

        # 查找表头行
        header_idx, header_row = self._find_header_row(all_rows)
        if header_idx < 0:
            return None

        # 数据行
        data_rows = []
        for row in all_rows[header_idx + 1:]:
            cleaned = [str(v) if v is not None else "" for v in row]
            # 跳过全空行
            if any(v.strip() for v in cleaned):
                data_rows.append(cleaned)

        if not data_rows:
            return None

        # 映射表头
        mapping = self._mapper.map_header(header_row)
        mapping = self._mapper.resolve_ambiguity(mapping, data_rows)

        # 分类表格
        classification = self._classifier.classify(header_row, data_rows)
        classification.header_mapping = mapping

        if not classification.is_field_table:
            return None

        # 提取视图标识
        view_code, view_name = self._extract_view_id(sheet_name, all_rows[:header_idx])

        # 提取字段
        fields = self._extract_fields(data_rows, mapping)

        if not fields:
            return None

        return ParsedView(
            view_code=view_code,
            view_name=view_name,
            fields=fields,
        )

    def _find_header_row(self, all_rows: List[Tuple]) -> Tuple[int, List[str]]:
        """查找表头行

        策略：从前几行中查找包含核心字段关键词的行
        """
        for i, row in enumerate(all_rows[:5]):
            cleaned = [str(v) if v is not None else "" for v in row]
            row_text = " ".join(cleaned)

            # 检查是否包含字段定义表的关键词
            has_field_keyword = any(kw in row_text for kw in ["字段", "名称", "数据类型", "数据元"])
            has_type_keyword = any(kw in row_text for kw in ["数据类型", "字段格式", "类型", "varchar", "VARCHAR"])

            if has_field_keyword and has_type_keyword:
                return i, cleaned

        return -1, []

    def _extract_view_id(self, sheet_name: str, title_rows: List[Tuple]) -> Tuple[str, str]:
        """从 Sheet 名和标题行提取视图标识

        Returns:
            (view_code, view_name)
        """
        view_code = ""
        view_name = ""

        # 从标题行提取视图代码
        for row in title_rows:
            row_text = " ".join(str(v) for v in row if v)
            # 模式1: 中文名（英文名）
            match = re.search(r'([一-龥]+)[（(]([A-Za-z_]\w+)[)）]', row_text)
            if match:
                view_name = match.group(1)
                view_code = match.group(2)
                return view_code, view_name

            # 模式2: 英文标识符
            match = re.search(r'([A-Z][A-Za-z_]\w+)', row_text)
            if match:
                view_code = match.group(1)

        # 从 Sheet 名提取
        if not view_code:
            # 模式: 数字 + 中文名
            match = re.match(r'^[\d]+[、.]?\s*(.+)', sheet_name)
            if match:
                view_name = match.group(1).strip()
            else:
                view_name = sheet_name

            # 尝试从第一行提取代码
            if title_rows:
                first_text = " ".join(str(v) for v in title_rows[0] if v)
                match = re.search(r'([A-Z][A-Za-z_]\w+)', first_text)
                if match:
                    view_code = match.group(1)

        if not view_code:
            view_code = view_name

        return view_code, view_name

    def _extract_fields(self, data_rows: List[List[str]], mapping: Dict[int, str]) -> List[ParsedField]:
        """从数据行提取字段定义"""
        fields = []

        # 找到各字段的列索引
        name_col = None
        code_col = None
        type_col = None
        length_col = None
        required_col = None
        desc_col = None

        for col_idx, field_type in mapping.items():
            if field_type == "field_name":
                code_col = col_idx
            elif field_type == "field_chinese_name":
                name_col = col_idx
            elif field_type == "data_type":
                type_col = col_idx
            elif field_type == "data_length":
                length_col = col_idx
            elif field_type == "required":
                required_col = col_idx
            elif field_type == "description":
                desc_col = col_idx

        # 至少需要 code 或 name
        if code_col is None and name_col is None:
            return []

        for row in data_rows:
            # 获取字段代码
            code = ""
            if code_col is not None and code_col < len(row):
                code = row[code_col].strip().replace(" ", "")  # 去掉所有空格

            # 获取字段中文名
            name = ""
            if name_col is not None and name_col < len(row):
                name = row[name_col].strip()

            # 跳过空行
            if not code and not name:
                continue

            # 跳过明显不是字段定义的行（如合计行）
            if code in ["合计", "总计", "小计"] or name in ["合计", "总计", "小计"]:
                continue

            # 获取数据类型
            field_type = ""
            if type_col is not None and type_col < len(row):
                field_type = row[type_col].strip()

            # 获取长度
            length = 0
            if length_col is not None and length_col < len(row):
                length_str = row[length_col].strip()
                # 提取数字
                match = re.search(r'\d+', length_str)
                if match:
                    length = int(match.group())

            # 获取是否必填
            required = False
            if required_col is not None and required_col < len(row):
                req_str = row[required_col].strip().upper()
                required = req_str in ["Y", "YES", "是", "必填", "非空", "TRUE", "1"]

            # 获取备注
            description = ""
            if desc_col is not None and desc_col < len(row):
                description = row[desc_col].strip()

            fields.append(ParsedField(
                name=name,
                code=code,
                field_type=field_type,
                length=length,
                required=required,
                description=description,
            ))

        return fields


def parse_xlsx(file_path: str, config_path: Optional[str] = None) -> ParseResult:
    """便捷函数：解析 XLSX 文件"""
    parser = XlsxParser(config_path)
    return parser.parse(file_path)
