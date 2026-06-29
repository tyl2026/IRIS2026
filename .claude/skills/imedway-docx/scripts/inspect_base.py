"""检查 base.docx 里的样式、numbering、header/footer 是否完整保留。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

from pathlib import Path
from docx import Document

BASE = Path(__file__).parent.parent / "assets" / "base.docx"
doc = Document(str(BASE))

print("=" * 60)
print("STYLES (按名称)")
print("=" * 60)
expected = [
    "5 文档标题", "1 标题 1", "1 标题 2", "1 标题 3", "1 标题 4", "1 标题 5",
    "2 正文", "3 标题 6", "3 标题 7", "3 标题 8", "3 标题 9",
    "4 表格表头", "4 表格内容", "4 表格标题", "声明日期", "Normal", "Normal Indent",
]
present = {s.name for s in doc.styles}
for name in expected:
    mark = "OK " if name in present else "MISS"
    print(f"  [{mark}] {name}")

print()
print("=" * 60)
print("段落 (前 20 段，看模板自带哪些示例)")
print("=" * 60)
for i, p in enumerate(doc.paragraphs[:20]):
    text = (p.text or "").replace("\n", " ")[:50]
    print(f"  [{i:2d}] style={p.style.name!r:15s} text={text!r}")

print()
print("=" * 60)
print("HEADER")
print("=" * 60)
section = doc.sections[0]
header = section.header
for i, p in enumerate(header.paragraphs):
    text = (p.text or "").replace("\n", " ")[:50]
    print(f"  [{i}] style={p.style.name!r:15s} text={text!r}")

# 看 header xml 里有没有图片关系
import re
xml = header.part.element.xml
n_drawing = len(re.findall(r"<w:drawing", xml))
n_pict = len(re.findall(r"<w:pict", xml))
n_blip = len(re.findall(r"<a:blip ", xml))
print(f"  drawings={n_drawing}, picts={n_pict}, blip(image refs)={n_blip}")

print()
print("=" * 60)
print("FOOTER")
print("=" * 60)
footer = section.footer
for i, p in enumerate(footer.paragraphs):
    text = (p.text or "").replace("\n", " ")[:80]
    print(f"  [{i}] style={p.style.name!r:15s} text={text!r}")

xml = footer.part.element.xml
print(f"  drawings={len(re.findall(r'<w:drawing', xml))}, picts={len(re.findall(r'<w:pict', xml))}")

print()
print("=" * 60)
print("NUMBERING (numId 列表)")
print("=" * 60)
try:
    numbering_part = doc.part.numbering_part
    num_xml = numbering_part.element.xml
    num_ids = sorted(set(re.findall(r'<w:num\s+w:numId="(\d+)"', num_xml)), key=int)
    abs_ids = sorted(set(re.findall(r'<w:abstractNum\s+w:abstractNumId="(\d+)"', num_xml)), key=int)
    print(f"  numIds: {num_ids}")
    print(f"  abstractNumIds: {abs_ids}")
except Exception as e:
    print(f"  [ERR] {e}")

print()
print("=" * 60)
print("PAGE SETUP")
print("=" * 60)
sec = doc.sections[0]
from docx.shared import Emu
def to_in(v):
    return f"{v.inches:.2f}in" if v else "None"
print(f"  page: {to_in(sec.page_width)} x {to_in(sec.page_height)}")
print(f"  margin L/R={to_in(sec.left_margin)}/{to_in(sec.right_margin)}, T/B={to_in(sec.top_margin)}/{to_in(sec.bottom_margin)}")
print(f"  header_distance={to_in(sec.header_distance)}, footer_distance={to_in(sec.footer_distance)}")
