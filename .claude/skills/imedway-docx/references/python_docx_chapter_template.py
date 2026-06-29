"""
iMEDWAY 章节版模板（含封面+目录+分节页码）— Base-Replace 模式。

基于 base_chapter.docx（章节版模板），提供：
- 封面页（标题+署名+日期）
- 自动目录（TOC 域，打开时更新）
- 正文分节（页码从1重新开始）
- 页眉占位符替换
- 页眉页脚继承（含模板 slogan 图片）

页脚处理策略：
  python-docx 无法可靠地在 footer 中持久化 XML 追加的图片（save 序列化会丢失）。
  因此：
  1. 页码文字用 python-docx API 写入（可持久化）
  2. slogan 图片用 python-docx 的 run.add_picture() 写入（可持久化）
  3. 目录节 footer 需在 save 后通过 ZIP 后处理彻底移除 footerReference
     （LibreOffice 不尊重空 footer，仍会渲染页码；必须移除引用本身）

用法:
    from python_docx_chapter_template import ImedwayChapterDoc

    doc = ImedwayChapterDoc(
        output_path="output.docx",
        title_lines=["<封面标题第一行>", "<封面标题第二行（可空）>"],
        author="<署名>",
        date="<日期>",
        header_name="<文档名>",
    )

    doc.add_heading_1("一级标题")
    doc.add_heading_2("二级标题")
    doc.add_heading_3("三级标题")
    doc.add_body("正文段落...")
    doc.add_bullet("要点 A", level=1)
    doc.add_table(headers=["列1","列2"], rows=[["a","b"]])
    doc.save()
"""
from __future__ import annotations

import shutil
import tempfile
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

# ── 模板路径 ────────────────────────────────────────────────────
_ASSETS = Path(__file__).parent.parent / "assets"
BASE_CHAPTER_DOCX = _ASSETS / "base_chapter.docx"
SLOGAN_PNG = _ASSETS / "slogan.png"
HEADER_PLACEHOLDER = "这里放置文档名称"

# ── 模板样式名（章节版，基于 Heading 1/2/3 继承编号）────────────
STYLE_H1 = "1 标题 1"       # 第X章
STYLE_H2 = "1 标题 2"       # 1.1.
STYLE_H3 = "1 标题 3"       # 1.1.1.
STYLE_BODY = "2 正文"       # 首行缩进 2 字符
STYLE_BULLET_1 = "3 点标题 1"  # ■
STYLE_BULLET_2 = "3 点标题 2"  # 子内容
STYLE_BULLET_3 = "3 点标题 3"  # ●
STYLE_BULLET_4 = "3 点标题 4"  # ○
STYLE_BULLET_5 = "3 点标题 5"
STYLE_BULLET_6 = "3 点标题 6"

# ── 封面结构（base_chapter.docx 段落索引）────────────────────────
COVER_TITLE_LINE1_IDX = 3
COVER_TITLE_LINE2_IDX = 4
COVER_COMPANY_IDX = 16
COVER_AUTHOR_IDX = 17
COVER_DATE_IDX = 18
COVER_LAST_IDX = 20  # 从此段开始清除

# slogan 图片原始尺寸（来自模板 footer4.xml 的 extent）
_SLOGAN_CX = 2300605
_SLOGAN_CY = 129540


# ── 工具函数 ─────────────────────────────────────────────────

def _set_cover_text(paragraph, new_text: str) -> None:
    """替换封面段落文本，保留第一个 run 的格式，清空其余 run。"""
    runs = paragraph.runs
    if not runs:
        return
    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def _add_table_borders(table) -> None:
    """给表格加完整黑色单线边框。"""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)


def _add_field(paragraph, field_code: str) -> None:
    """在段落中插入一个 Word 域（如 PAGE、NUMPAGES）。"""
    run1 = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run1._element.append(fld_begin)

    run2 = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {field_code} "
    run2._element.append(instr)

    run3 = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run3._element.append(fld_sep)

    paragraph.add_run("1")

    run5 = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run5._element.append(fld_end)


def _remove_footer_refs_from_toc(docx_path: Path) -> None:
    """ZIP 后处理：移除目录节的所有 footerReference。

    LibreOffice 不尊重空 footer 内容——即使 footer XML 里只有空段落，
    仍会渲染页码。唯一可靠的方式是从 sectPr 中彻底移除 footerReference，
    让目录节不关联任何 footer part。

    识别策略：目录节的 default footer 被清空了（无文字无图片），
    而封面和正文的 footer 有实际内容。找到 default footer 为空的 sectPr，
    移除其全部 footerReference。
    """
    from lxml import etree

    ns_w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

    tmp_path = docx_path.with_suffix('.tmp')

    def _footer_is_empty(zin, rels_map, sectPr) -> bool:
        """检查 sectPr 的 default footer 是否为空（无文字无图片）。"""
        for ref in sectPr.findall(f'{{{ns_w}}}footerReference'):
            if ref.get(f'{{{ns_w}}}type') == 'default':
                rid = ref.get(f'{{{ns_r}}}id')
                target = rels_map.get(rid)
                if not target:
                    return True
                try:
                    fc = zin.read(f'word/{target}')
                except KeyError:
                    return True
                fr = etree.fromstring(fc)
                for r in fr.iter(f'{{{ns_w}}}t'):
                    if r.text and r.text.strip():
                        return False
                if fr.find(f'.//{{{ns_w}}}drawing') is not None:
                    return False
                return True
        return True  # no default footer ref at all

    with zipfile.ZipFile(docx_path, 'r') as zin:
        # Build rId -> Target map from document.xml.rels
        rels_root = etree.fromstring(zin.read('word/_rels/document.xml.rels'))
        rels_map = {rel.get('Id'): rel.get('Target') for rel in rels_root}

        with zipfile.ZipFile(tmp_path, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == 'word/document.xml':
                    root = etree.fromstring(data)
                    body = root.find(f'{{{ns_w}}}body')

                    # Collect all sectPrs (inline first, then body-end)
                    sectPrs = []
                    for child in body:
                        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                        if tag == 'p':
                            pPr = child.find(f'{{{ns_w}}}pPr')
                            if pPr is not None:
                                secPr = pPr.find(f'{{{ns_w}}}sectPr')
                                if secPr is not None:
                                    sectPrs.append(secPr)
                    body_sectPr = body.find(f'{{{ns_w}}}sectPr')
                    if body_sectPr is not None:
                        sectPrs.append(body_sectPr)

                    # Strategy: remove footer refs from TOC section AND cover section's
                    # default/even footers (LibreOffice inherits from cover → TOC if TOC has none).
                    # Cover has titlePg, so only its "first" footer is used (which is empty).
                    modified = False
                    for idx, sp in enumerate(sectPrs):
                        refs = sp.findall(f'{{{ns_w}}}footerReference')
                        if not refs:
                            continue

                        if idx == 0:
                            # Section 0 (cover): remove default+even, keep first (empty)
                            for ref in refs:
                                rtype = ref.get(f'{{{ns_w}}}type')
                                if rtype in ('default', 'even'):
                                    sp.remove(ref)
                                    modified = True
                        elif _footer_is_empty(zin, rels_map, sp):
                            # TOC section: remove all
                            for ref in refs:
                                sp.remove(ref)
                            modified = True

                    if modified:
                        data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)

                zout.writestr(item, data)

    tmp_path.replace(docx_path)


# ── 主类 ─────────────────────────────────────────────────────

class ImedwayChapterDoc:
    """iMEDWAY 章节版模板包装。含封面+目录+分节页码。"""

    def __init__(
        self,
        output_path: str | Path,
        title_lines: list[str] | None = None,
        author: str = "XXX部门",
        date: str = "XXXX年X月X日",
        company: str = "东华医为科技有限公司",
        header_name: str = HEADER_PLACEHOLDER,
        base_docx: str | Path | None = None,
    ):
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        # 1. 复制章节版 base
        base = Path(base_docx) if base_docx else BASE_CHAPTER_DOCX
        if not base.exists():
            raise FileNotFoundError(f"base_chapter.docx missing: {base}")
        shutil.copy(base, self.output_path)

        # 2. 打开
        self.doc = Document(str(self.output_path))

        # 3. 修改封面
        paras = self.doc.paragraphs
        if title_lines:
            _set_cover_text(paras[COVER_TITLE_LINE1_IDX], title_lines[0])
            _set_cover_text(paras[COVER_TITLE_LINE2_IDX],
                            title_lines[1] if len(title_lines) >= 2 else "")
        _set_cover_text(paras[COVER_COMPANY_IDX], company)
        _set_cover_text(paras[COVER_AUTHOR_IDX], author)
        _set_cover_text(paras[COVER_DATE_IDX], date)

        # 4. 替换页眉占位符
        if header_name != HEADER_PLACEHOLDER:
            for section in self.doc.sections:
                for paragraph in section.header.paragraphs:
                    if HEADER_PLACEHOLDER in paragraph.text:
                        for run in paragraph.runs:
                            if HEADER_PLACEHOLDER in run.text:
                                run.text = run.text.replace(
                                    HEADER_PLACEHOLDER, header_name
                                )

        # 5. 清空版本历史页+TOC+正文（保留 sectPr）
        body = self.doc.element.body
        sectPr = body.find(qn("w:sectPr"))
        children = list(body)
        count = 0
        for child in children:
            if child is sectPr:
                continue
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag in ('p', 'tbl'):
                count += 1
                if count > COVER_LAST_IDX:
                    body.remove(child)

        # 6. 插入目录
        self._add_toc()

        # 7. 文档属性
        if title_lines:
            self.doc.core_properties.title = " ".join(title_lines)
        self.doc.core_properties.author = author

    def _add_toc(self) -> None:
        """插入 TOC 域 + 正文分节（页码从1开始）。"""
        doc = self.doc

        # 设置打开时自动更新域
        settings = doc.settings.element
        update_fields = settings.find(qn("w:updateFields"))
        if update_fields is None:
            update_fields = OxmlElement("w:updateFields")
            settings.append(update_fields)
        update_fields.set(qn("w:val"), "true")

        # 目录标题
        p_title = doc.add_paragraph()
        p_title.add_run().add_break()
        p_title.add_run().add_break()
        p_heading = doc.add_paragraph()
        p_heading.alignment = 1  # center
        rh = p_heading.add_run("目 录")
        rh.bold = True
        rh.font.size = Pt(22)

        # TOC 域（显式指定自定义样式名）
        paragraph = doc.add_paragraph()
        run = paragraph.add_run()
        fldChar1 = OxmlElement("w:fldChar")
        fldChar1.set(qn("w:fldCharType"), "begin")
        run._element.append(fldChar1)

        run2 = paragraph.add_run()
        instrText = OxmlElement("w:instrText")
        instrText.set(qn("xml:space"), "preserve")
        instrText.text = (
            ' TOC \\o "1-3" \\h \\z \\u '
            '\\t "1 标题 1,1,1 标题 2,2,1 标题 3,3" '
        )
        run2._element.append(instrText)

        run3 = paragraph.add_run()
        fldChar2 = OxmlElement("w:fldChar")
        fldChar2.set(qn("w:fldCharType"), "separate")
        run3._element.append(fldChar2)

        paragraph.add_run()

        run5 = paragraph.add_run()
        fldChar3 = OxmlElement("w:fldChar")
        fldChar3.set(qn("w:fldCharType"), "end")
        run5._element.append(fldChar3)

        # 创建正文节 + 配置页脚
        toc_section = doc.sections[1]
        self._setup_body_section(toc_section)

    def _setup_body_section(self, toc_section) -> None:
        """清空目录 footer → 创建正文节 → python-docx API 写页码 + slogan 图片。"""
        doc = self.doc

        # 1. 保存模板 default footer 的 P0 格式（在清空之前）
        sectPr1 = toc_section._sectPr
        orig_pPr_xml = None
        for ref in sectPr1.findall(qn('w:footerReference')):
            if ref.get(qn('w:type')) == 'default':
                orig_rId = ref.get(qn('r:id'))
                orig_fp = toc_section.part.rels[orig_rId].target_part
                paras = orig_fp._element.findall(qn('w:p'))
                if paras:
                    pPr = paras[0].find(qn('w:pPr'))
                    if pPr is not None:
                        orig_pPr_xml = deepcopy(pPr)
                break

        # 2. 清空目录节 footer（必须在 add_section 之前）
        toc_section.different_first_page_header_footer = False
        for p in toc_section.footer.paragraphs:
            for child in list(p._element):
                tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
                if tag != 'pPr':
                    p._element.remove(child)

        # 3. 创建正文节
        new_section = doc.add_section()
        new_section.start_type = 2  # NEW_PAGE

        # 4. 设置 pgNumType start=1（body-end sectPr + inline sectPr）
        for target_sectPr in self._get_all_sectPrs():
            pnt = target_sectPr.find(qn('w:pgNumType'))
            if pnt is None:
                pnt = OxmlElement('w:pgNumType')
                target_sectPr.append(pnt)
            pnt.set(qn('w:start'), '1')

        new_section.header.is_linked_to_previous = True
        new_section.footer.is_linked_to_previous = False
        new_section.different_first_page_header_footer = False

        # 5. 用 python-docx API 写页码（可持久化）
        p0 = new_section.footer.paragraphs[0]
        for child in list(p0._element):
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag != 'pPr':
                p0._element.remove(child)

        p0.add_run('第')
        _add_field(p0, r'PAGE \* Arabic \* MERGEFORMAT')
        p0.add_run('页')

        # 应用模板 P0 的段落格式（居中、楷体 9pt、spacing after=120）
        if orig_pPr_xml is not None:
            existing_pPr = p0._element.find(qn('w:pPr'))
            if existing_pPr is not None:
                p0._element.remove(existing_pPr)
            p0._element.insert(0, deepcopy(orig_pPr_xml))

        # 6. 添加 slogan 图片（python-docx API，可持久化）
        if SLOGAN_PNG.exists():
            p_img = new_section.footer.add_paragraph()
            p_img.alignment = 1  # center
            run_img = p_img.add_run()
            run_img.add_picture(str(SLOGAN_PNG), width=Emu(_SLOGAN_CX))

    def _get_all_sectPrs(self) -> list:
        """获取文档中所有 sectPr 元素（inline + body-end）。

        python-docx add_section() 的 sectPr 分两部分：
          - body 末尾的 sectPr（最后 section 的属性）
          - 段落中的 inline sectPr（分节符）
        返回 [Section0_inline, Section1_inline, ..., SectionN_body_end]
        """
        body = self.doc.element.body
        sectPrs = []
        for child in list(body):
            tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
            if tag == 'p':
                pPr = child.find(qn("w:pPr"))
                if pPr is not None:
                    secPr = pPr.find(qn("w:sectPr"))
                    if secPr is not None:
                        sectPrs.append(secPr)
        body_sectPr = body.find(qn("w:sectPr"))
        if body_sectPr is not None:
            sectPrs.append(body_sectPr)
        return sectPrs

    # ── 段落 API ───────────────────────────────────────────

    def add_heading_1(self, text: str):
        """一级标题（第X章）。"""
        return self.doc.add_paragraph(text, style=STYLE_H1)

    def add_heading_2(self, text: str):
        """二级标题（X.Y.）。"""
        return self.doc.add_paragraph(text, style=STYLE_H2)

    def add_heading_3(self, text: str):
        """三级标题（X.Y.Z.）。"""
        return self.doc.add_paragraph(text, style=STYLE_H3)

    def add_body(self, text: str):
        """正文段落（首行缩进 2 字符）。"""
        return self.doc.add_paragraph(text, style=STYLE_BODY)

    def add_bullet(self, text: str, level: int = 1):
        """项目符号。level: 1-6 对应 3 点标题 1-6。"""
        styles = {
            1: STYLE_BULLET_1, 2: STYLE_BULLET_2,
            3: STYLE_BULLET_3, 4: STYLE_BULLET_4,
            5: STYLE_BULLET_5, 6: STYLE_BULLET_6,
        }
        style = styles.get(level, STYLE_BULLET_1)
        return self.doc.add_paragraph(text, style=style)

    # ── 表格 API ───────────────────────────────────────────

    def add_table(self, headers: list[str], rows: list[list[str]]):
        """标准表格（表头加粗，单线边框）。"""
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(h))
            run.bold = True
        for r, row_data in enumerate(rows):
            for c, val in enumerate(row_data):
                cell = table.rows[r + 1].cells[c]
                cell.text = ""
                cell.paragraphs[0].add_run(str(val))
        _add_table_borders(table)
        return table

    # ── 保存 ──────────────────────────────────────────────

    def save(self, output_path: str | Path | None = None):
        """保存文档。先 python-docx save，再 ZIP 后处理移除目录节 footer。"""
        path = Path(output_path) if output_path else self.output_path
        self.doc.save(str(path))
        _remove_footer_refs_from_toc(path)
        return path
