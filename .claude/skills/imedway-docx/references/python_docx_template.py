"""
iMEDWAY (东华医为) 小页眉 DOCX 模板包装 — Base-Replace 模式。

不从零重建样式/页眉/页脚/编号定义；而是把官方模板 base.docx 复制一份，清空 body
（保留 sectPr → 保留页面尺寸 + header/footer 引用），然后用模板自带的样式名
("5 文档标题"/"1 标题 1"/"2 正文" 等) 追加内容。编号要靠段落级 numPr 显式挂载，
这是 LibreOffice .doc→.docx 转换后必须手动激活的关键步骤。

用法:
    from python_docx_template import ImedwayDoc

    doc = ImedwayDoc(
        output_path="output.docx",
        header_name="<文档名>",
        title="<文档名>",
        author="<署名>",
    )
    doc.add_heading_1("引言")
    doc.add_heading_2("调研背景")
    doc.add_body("正文段落...")
    doc.add_bracket_heading("赛道分布")
    doc.add_sub_content("子内容...")
    doc.add_bullet("要点 A", level=1)
    doc.add_table(headers=["列1","列2"], rows=[["a","b"]], caption="样例表")
    doc.add_closing(dept="<署名>", date="<日期>")
    doc.save()
"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# ── 模板自带样式名 ─────────────────────────────────────────────
STYLE_TITLE       = "5 文档标题"   # 22pt Bold Centered
STYLE_H1          = "1 标题 1"     # 1.
STYLE_H2          = "1 标题 2"     # 1.1.
STYLE_H3          = "1 标题 3"     # 1.1.1.
STYLE_BRACKET     = "1 标题 4"     # 1）
STYLE_PAREN       = "1 标题 5"     # (1)
STYLE_BODY        = "2 正文"       # 首行缩进 2 字符
STYLE_BULLET_1    = "3 标题 6"     # ■
STYLE_BULLET_2    = "3 标题 7"     # □
STYLE_BULLET_3    = "3 标题 8"     # ●
STYLE_BULLET_4    = "3 标题 9"     # ○
STYLE_TABLE_HEAD  = "4 表格表头"
STYLE_TABLE_BODY  = "4 表格内容"
STYLE_TABLE_TITLE = "4 表格标题"
STYLE_CLOSING     = "声明日期"     # 14pt Bold Centered

# ── 段落级 numPr 映射（base.docx 中 LibreOffice 分配的 numId） ─────
# 这些 numId 必须在段落级显式挂载，否则 Word/LibreOffice 不会渲染编号
NUMBERING = {
    STYLE_H1:       (5, 0),
    STYLE_H2:       (5, 1),
    STYLE_H3:       (5, 2),
    STYLE_BRACKET:  (3, 0),
    STYLE_PAREN:    (11, 0),
    STYLE_BULLET_1: (8, 1),
    STYLE_BULLET_2: (10, 1),
    STYLE_BULLET_3: (9, 1),
    STYLE_BULLET_4: (6, 1),
}

BASE_DOCX = Path(__file__).parent.parent / "assets" / "base.docx"
HEADER_PLACEHOLDER = "这里放置文档名称"


# ── 工具函数 ─────────────────────────────────────────────────

def _attach_numbering(paragraph, num_id: int, ilvl: int) -> None:
    """在段落级显式挂载 numPr，覆盖任何已有编号。"""
    pPr = paragraph._element.get_or_add_pPr()
    for old in pPr.findall(qn("w:numPr")):
        pPr.remove(old)
    numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(ilvl)); numPr.append(il)
    nid = OxmlElement("w:numId"); nid.set(qn("w:val"), str(num_id)); numPr.append(nid)
    pPr.insert(0, numPr)


def _clear_body(doc) -> None:
    """删除 body 中的所有段落和表格，仅保留 sectPr。"""
    body = doc.element.body
    sectPr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sectPr:
            body.remove(child)


def _replace_header_placeholder(doc, new_name: str) -> None:
    """把 header 里的 '这里放置文档名称' 占位文本替换为指定文档名。保留字体格式。"""
    if new_name == HEADER_PLACEHOLDER:
        return
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if HEADER_PLACEHOLDER not in paragraph.text:
                continue
            # 改 run 内的 w:t，逐个 run 处理，保留字体格式
            for run in paragraph.runs:
                if HEADER_PLACEHOLDER in run.text:
                    run.text = run.text.replace(HEADER_PLACEHOLDER, new_name)
                    return  # 占位符通常只在一个 run 内


def _add_table_borders(table) -> None:
    """给表格加完整黑色单线边框（与模板原表一致）。"""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    # 移除已有 tblBorders
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "000000")
        borders.append(b)
    tblPr.append(borders)


# ── 主类 ─────────────────────────────────────────────────────

class ImedwayDoc:
    """iMEDWAY 小页眉模板的薄包装。所有样式/页眉/页脚/编号来自 base.docx。"""

    def __init__(
        self,
        output_path: str | Path,
        header_name: str = HEADER_PLACEHOLDER,
        title: str | None = None,
        author: str | None = None,
        base_docx: str | Path | None = None,
    ):
        self.output_path = Path(output_path)
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        # 1. 复制 base.docx 作为输出底 — 保留 styles/numbering/header/footer/sectPr
        base = Path(base_docx) if base_docx else BASE_DOCX
        if not base.exists():
            raise FileNotFoundError(f"base.docx 缺失: {base}")
        shutil.copy(base, self.output_path)

        # 2. 打开后清空 body
        self.doc = Document(str(self.output_path))
        _clear_body(self.doc)

        # 3. 替换 header 占位符
        _replace_header_placeholder(self.doc, header_name)

        # 4. 文档属性
        if author:
            self.doc.core_properties.author = author
        if title:
            self.doc.core_properties.title = title
            self.add_title(title)

    # ── 段落 API ───────────────────────────────────────────

    def add_title(self, text: str):
        """文档标题 (5 文档标题，22pt Bold Centered)。"""
        return self.doc.add_paragraph(text, style=STYLE_TITLE)

    def add_heading_1(self, text: str):
        """一级标题，自动编号 '1.'。"""
        p = self.doc.add_paragraph(text, style=STYLE_H1)
        _attach_numbering(p, *NUMBERING[STYLE_H1])
        return p

    def add_heading_2(self, text: str):
        """二级标题，自动编号 '1.1.'。"""
        p = self.doc.add_paragraph(text, style=STYLE_H2)
        _attach_numbering(p, *NUMBERING[STYLE_H2])
        return p

    def add_heading_3(self, text: str):
        """三级标题，自动编号 '1.1.1.'。"""
        p = self.doc.add_paragraph(text, style=STYLE_H3)
        _attach_numbering(p, *NUMBERING[STYLE_H3])
        return p

    def add_bracket_heading(self, text: str):
        """单括号标题，自动编号 '1）'。"""
        p = self.doc.add_paragraph(text, style=STYLE_BRACKET)
        _attach_numbering(p, *NUMBERING[STYLE_BRACKET])
        return p

    def add_sub_content(self, text: str):
        """双括号子内容，自动编号 '(1)'。"""
        p = self.doc.add_paragraph(text, style=STYLE_PAREN)
        _attach_numbering(p, *NUMBERING[STYLE_PAREN])
        return p

    def add_body(self, text: str):
        """正文段落 (12pt 宋体，首行缩进 2 字符，1.25 倍行距)。"""
        return self.doc.add_paragraph(text, style=STYLE_BODY)

    def add_bullet(self, text: str, level: int = 1):
        """项目符号列表。level: 1=■, 2=□, 3=●, 4=○。"""
        style = {1: STYLE_BULLET_1, 2: STYLE_BULLET_2,
                 3: STYLE_BULLET_3, 4: STYLE_BULLET_4}.get(level, STYLE_BULLET_1)
        p = self.doc.add_paragraph(text, style=style)
        _attach_numbering(p, *NUMBERING[style])
        return p

    def add_blank(self):
        """空段落，用于段落间距分隔。"""
        return self.doc.add_paragraph("", style=STYLE_BODY)

    # ── 表格 API ───────────────────────────────────────────

    def add_table(self, headers: list[str], rows: list[list[str]], caption: str | None = None):
        """
        创建标准 iMEDWAY 表格。

        - caption: 表格标题 (会渲染为右对齐 '-标题-')
        - headers: 表头列表 (4 表格表头样式，宋体 Bold 居中)
        - rows: 数据行 (4 表格内容样式，宋体 左对齐)
        """
        if caption:
            cap = self.doc.add_paragraph(f"-{caption}-", style=STYLE_TABLE_TITLE)
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        # 表头
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = self.doc.styles[STYLE_TABLE_HEAD]
            p.add_run(str(h))
        # 数据行
        for r, row_data in enumerate(rows):
            for c, val in enumerate(row_data):
                cell = table.rows[r + 1].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                p.style = self.doc.styles[STYLE_TABLE_BODY]
                p.add_run(str(val))
        _add_table_borders(table)
        return table

    # ── 结尾区块 API ───────────────────────────────────────

    def add_closing(
        self,
        company: str = "东华医为科技有限公司",
        dept: str = "部门名称",
        date: str = "XXXX年X月X日",
    ):
        """文末三行 (公司名 / 部门 / 日期)，14pt Bold Centered。"""
        for text in (company, dept, date):
            self.doc.add_paragraph(text, style=STYLE_CLOSING)

    # ── 保存 ──────────────────────────────────────────────

    def save(self, output_path: str | Path | None = None):
        """保存文档。若指定新路径则保存到新位置，否则覆盖初始化时的路径。"""
        path = Path(output_path) if output_path else self.output_path
        self.doc.save(str(path))
        return path


# ── 冒烟测试 ─────────────────────────────────────────────────

if __name__ == "__main__":
    import sys, io
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

    out = Path(__file__).parent.parent / "build" / "_smoke_output.docx"

    doc = ImedwayDoc(
        output_path=out,
        header_name="<文档名>",
        title="<文档名>",
        author="<署名>",
    )

    doc.add_heading_1("一级标题示例")
    doc.add_heading_2("二级标题示例")
    doc.add_heading_3("三级标题示例")
    doc.add_body("正文段落示例。")

    doc.add_bracket_heading("括号标题示例")
    doc.add_sub_content("子内容示例。")
    doc.add_bullet("要点 A", level=1)
    doc.add_bullet("要点 B", level=2)

    doc.add_table(
        headers=["列1", "列2", "列3"],
        rows=[
            ["A1", "B1", "C1"],
            ["A2", "B2", "C2"],
        ],
        caption="示例表",
    )

    doc.add_closing(
        company="东华医为科技有限公司",
        dept="<署名>",
        date="<日期>",
    )

    saved = doc.save()
    print(f"OK saved: {saved}")
